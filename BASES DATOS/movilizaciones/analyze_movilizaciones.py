import os
from openpyxl import load_workbook

# Directory containing the Excel files
dir_path = os.path.dirname(os.path.abspath(__file__))

# List of files to analyze
files = [
    'movilizaciones_UGROY.xlsx',
    'movilizaciones_UGRY.xlsx',
]

def analyze_excel(file_path):
    print(f'\n=== {os.path.basename(file_path)} ===')
    wb = load_workbook(file_path, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f'\nSheet: {sheet_name}')
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            print('  (Empty sheet)')
            continue
        headers = rows[0]
        print('  Headers:', headers)
        for i, row in enumerate(rows[1:4], 1):
            print(f'  Row {i}:', row)

def main():
    for fname in files:
        fpath = os.path.join(dir_path, fname)
        if os.path.exists(fpath):
            analyze_excel(fpath)
        else:
            print(f'File not found: {fname}')

if __name__ == '__main__':
    main()
