import os
import json
from openpyxl import load_workbook
from datetime import datetime

dir_path = os.path.dirname(os.path.abspath(__file__))
files = [
    'movilizaciones_UGROY.xlsx',
    'movilizaciones_UGRY.xlsx',
]

def excel_date(dt):
    if isinstance(dt, datetime):
        return dt.strftime('%Y-%m-%d')
    if isinstance(dt, str):
        return dt
    return ''

# Recursively convert all datetime objects in data to strings
def convert_datetimes(obj):
    if isinstance(obj, dict):
        return {k: convert_datetimes(v) for k, v in obj.items()}
    elif isinstance(obj, list):
        return [convert_datetimes(x) for x in obj]
    elif isinstance(obj, datetime):
        return obj.strftime('%Y-%m-%d')
    else:
        return obj

def parse_matrix_sheet(ws, id_col=0, month_cols=None, total_col=None, label='ventanilla'):
    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return []
    headers = rows[0]
    data = []
    for row in rows[1:]:
        if not row or not row[id_col]:
            continue
        obj = {label: row[id_col]}
        if month_cols:
            obj['monthly'] = {}
            for idx, col in month_cols:
                val = row[idx]
                if val is not None:
                    obj['monthly'][excel_date(col)[:7]] = val
        if total_col is not None:
            obj['acumulado'] = row[total_col]
        data.append(obj)
    return data

def parse_inventario_especies(ws):
    rows = list(ws.iter_rows(values_only=True))
    data = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        data.append({
            'especie': row[0],
            'altas': row[1],
            'inventario_altas': row[2],
            'actualizaciones': row[3],
            'inventario_act': row[4],
        })
    return data

def parse_atencion_upp(ws):
    rows = list(ws.iter_rows(values_only=True))
    data = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        data.append({
            'especie': row[0],
            'upp_vigentes': row[1],
            'upp_suspendidas': row[2],
            'upp_actualizadas': row[3],
            'avance': row[4],
        })
    return data

def parse_identificadores(ws):
    rows = list(ws.iter_rows(values_only=True))
    data = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        data.append({
            'mes': excel_date(row[0])[:7],
            'upp': row[1],
            'solicitados': row[2],
            'entregados': row[3],
            'devueltos': row[4],
            'capturados': row[5],
            'avance': row[6],
        })
    return data

def parse_prestadores(ws):
    rows = list(ws.iter_rows(values_only=True))
    data = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        data.append({
            'actividad': row[0],
            'tipo': row[3],
            'historico': row[4],
            'opinion_positiva': row[5],
            'proceso_concluido': row[6],
            'nuevo_ingreso': row[7],
        })
    return data

def parse_rastro_monthly(ws, id_col=0, month_cols=None, total_col=None, label='rastro'):
    return parse_matrix_sheet(ws, id_col, month_cols, total_col, label)

def parse_identificadores_disponibles(ws):
    rows = list(ws.iter_rows(values_only=True))
    data = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        obj = {'ventanilla': row[0]}
        for idx, val in enumerate(row[1:], 1):
            if val is not None:
                obj[f'col{idx}'] = val
        data.append(obj)
    return data

def parse_informe_individua(ws):
    rows = list(ws.iter_rows(values_only=True))
    data = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        obj = {'clave': row[0]}
        for idx, val in enumerate(row[1:], 1):
            obj[f'col{idx}'] = val
        data.append(obj)
    return data

def parse_informe_reemo(ws):
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]
    data = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        obj = {'destino': row[0], 'guias': {}}
        for idx, val in enumerate(row[1:13], 1):
            if val is not None:
                month = headers[idx]
                if isinstance(month, str):
                    obj['guias'][month] = val
        data.append(obj)
    return data

def normalize_sheet_name(sheet):
    return sheet.strip().lower().replace(' ', '_').replace('.', '').replace('/', '_')

def main():
    for fname in files:
        fpath = os.path.join(dir_path, fname)
        if not os.path.exists(fpath):
            print(f'File not found: {fname}')
            continue
        print(f'\n=== {fname} ===')
        wb = load_workbook(fpath, data_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            print(f'\nSheet: {sheet}')
            # Sheet-specific logic
            if sheet.lower().startswith('altas de upp') or sheet.lower().startswith('altas de psg') or sheet.lower().startswith('actualizaciones de upp'):
                # Find month columns and acumulado
                headers = list(ws.iter_rows(values_only=True))[0]
                month_cols = [(i, h) for i, h in enumerate(headers) if isinstance(h, datetime)]
                total_col = None
                for i, h in enumerate(headers):
                    if isinstance(h, str) and 'acumulado' in h.lower():
                        total_col = i
                data = parse_matrix_sheet(ws, id_col=0, month_cols=month_cols, total_col=total_col)
            elif sheet.lower().startswith('inventario otras especies'):
                data = parse_inventario_especies(ws)
            elif sheet.lower().startswith('atencion upp por especie'):
                data = parse_atencion_upp(ws)
            elif sheet.lower().startswith('identificadores capturados') or sheet.lower().startswith('barrido'):
                data = parse_identificadores(ws)
            elif sheet.lower().startswith('prestadores de servicios'):
                data = parse_prestadores(ws)
            elif sheet.lower().startswith('identifiacores recuperados en r') or sheet.lower().startswith('identific. ord. matanza docs co'):
                headers = list(ws.iter_rows(values_only=True))[0]
                month_cols = [(i, h) for i, h in enumerate(headers) if isinstance(h, datetime)]
                total_col = None
                for i, h in enumerate(headers):
                    if isinstance(h, str) and 'total' in h.lower():
                        total_col = i
                data = parse_rastro_monthly(ws, id_col=0, month_cols=month_cols, total_col=total_col, label='rastro')
            elif sheet.lower().startswith('identificadores disponibles'):
                data = parse_identificadores_disponibles(ws)
            elif sheet.lower().startswith('informe individua'):
                data = parse_informe_individua(ws)
            elif sheet.lower().startswith('informe reemo'):
                data = parse_informe_reemo(ws)
            elif sheet.lower().startswith('registros de pgn'):
                # Custom: parse as vertical table
                rows = list(ws.iter_rows(values_only=True))
                data = []
                for row in rows[2:]:
                    if not row or not row[0]:
                        continue
                    obj = {
                        'fecha': excel_date(row[0]),
                        'tipo': row[1],
                        'UPP': row[2],
                        'PSG': row[3],
                        'PG_Equidos': row[4],
                        'PG_Ganado': row[5],
                    }
                    data.append(obj)
            else:
                print('  (No parser for this sheet)')
                continue
            # Save to JSON file
            out_name = f"{os.path.splitext(fname)[0]}__{normalize_sheet_name(sheet)}.json"
            out_path = os.path.join(dir_path, out_name)
            with open(out_path, 'w', encoding='utf-8') as f:
                json.dump(convert_datetimes(data), f, ensure_ascii=False, indent=2)
            print(f"  Saved: {out_name} ({len(data)} rows)")
            # Print preview
            print(json.dumps(convert_datetimes(data[:3]), ensure_ascii=False, indent=2))

if __name__ == '__main__':
    main()
